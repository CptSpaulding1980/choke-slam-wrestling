#!/usr/bin/env python3
"""Persistent post-processing for the generated CSW website.

The legacy match generator remains authoritative for XML-derived match and title
history. This module adds data-driven team rosters, champion profile panels,
Season 8 videos, media preservation, and event-link validation.
"""

from __future__ import annotations

import argparse
import html
import json
import re
import sys
import xml.etree.ElementTree as ET
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import Dict, Iterable, List, Mapping, Sequence, Set, Tuple

import yaml
from openpyxl import load_workbook


DEFAULT_SHEET_ID = "12IlPkUuXUVPZj7fSz_HmyaLmjhZWLY0C"
DEFAULT_SHEET_GID = 828908512
DEFAULT_CREDENTIALS = Path(
    "/Users/homeserver/.hermes/profiles/hermes-wrestling/data/google-key.json"
)
TEAM_SHEET_NAME = "Choke Slam Wrestling S8"
CACHE_RELATIVE_PATH = Path("data/team-rosters.json")

MANAGER_TO_TEAM = {
    "Pascal": "Saint Rebel Radicalz",
    "Phillipp": "Militanter Mummenschanz",
    "Philipp": "Militanter Mummenschanz",
    "Hendrik": "Sweet 'n Sour Elite",
}
TEAM_CONFIG = {
    "Militanter Mummenschanz": {
        "manager": "Philipp Brunkovic",
        "manager_slug": "philipp-brunkovic",
        "manager_image": "Philipp_Brunkovic.png",
        "logo": "ChokeSlam_MM.png",
        "audio": "Militanter_Mummenschanz.mp3",
        "slug": "militanter-mummenschanz",
    },
    "Saint Rebel Radicalz": {
        "manager": "Pascal LePas",
        "manager_slug": "pascal-le-pas",
        "manager_image": "Pascal_LePas.png",
        "logo": "ChokeSlam_SRR.png",
        "audio": "Saint_Rebel_Radicalz.mp3",
        "slug": "saint-rebel-radicalz",
    },
    "Sweet 'n Sour Elite": {
        "manager": "Hendrique Delafuente",
        "manager_slug": "hendrique-delafuente",
        "manager_image": "Hendrique_Delafuente.png",
        "logo": "ChokeSlam_SnS.png",
        "audio": "Sweet_n_Sour_Elite.mp3",
        "slug": "sweet-n-sour-elite",
    },
}
NAME_ALIASES = {
    "mankind": "Mick Foley",
    "sanada": "SANADA",
    "konosuke takeshita": "Konosuke Takeshita",
}
CHAMPION_START = "<!-- AUTO CHAMPION STATUS START -->"
CHAMPION_END = "<!-- AUTO CHAMPION STATUS END -->"
HIGHLIGHTS_START = "<!-- AUTO S08 HIGHLIGHTS START -->"
HIGHLIGHTS_END = "<!-- AUTO S08 HIGHLIGHTS END -->"
TEAM_START = "<!-- AUTO TEAM ROSTER START -->"
TEAM_END = "<!-- AUTO TEAM ROSTER END -->"


@dataclass(frozen=True)
class TeamSheetResult:
    tab_name: str
    colour_to_manager: Mapping[str, str]
    teams: Mapping[str, List[str]]


@dataclass(frozen=True)
class ValidationResult:
    checked: int
    missing: Sequence[str]


def parse_frontmatter(text: str) -> Tuple[dict, str]:
    """Parse only a leading YAML frontmatter block, without blind splitting."""
    lines = text.splitlines(keepends=True)
    if not lines or lines[0].strip() != "---":
        return {}, text
    closing = next((index for index in range(1, len(lines)) if lines[index].strip() == "---"), None)
    if closing is None:
        raise ValueError("Unterminated YAML frontmatter")
    data = yaml.safe_load("".join(lines[1:closing])) or {}
    if not isinstance(data, dict):
        raise ValueError("YAML frontmatter must be a mapping")
    return data, "".join(lines[closing + 1 :])


def _theme_colours(workbook) -> List[str]:
    if not workbook.loaded_theme:
        return []
    root = ET.fromstring(workbook.loaded_theme)
    namespace = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    scheme = root.find(".//a:clrScheme", namespace)
    if scheme is None:
        return []
    colours = []
    for colour in scheme:
        value = next(iter(colour), None)
        colours.append((value.get("lastClr") or value.get("val")) if value is not None else "")
    return colours


def _colour_hex(colour, workbook) -> str:
    value = ""
    if colour.type == "rgb" and colour.rgb:
        value = colour.rgb
    elif colour.type == "theme" and colour.theme is not None:
        theme = _theme_colours(workbook)
        if colour.theme < len(theme):
            value = theme[colour.theme]
    elif colour.type == "indexed" and colour.indexed is not None:
        from openpyxl.styles.colors import COLOR_INDEX

        if colour.indexed < len(COLOR_INDEX):
            value = COLOR_INDEX[colour.indexed]
    value = str(value).upper()
    if len(value) == 8:
        value = value[2:]
    return f"#{value}" if len(value) == 6 else ""


def _clean_sheet_name(value: object) -> str:
    name = re.sub(r"\s+\d+\)\s*$", "", str(value or "")).strip()
    return re.sub(r"\s+", " ", name)


def _resolve_name(raw_name: str, available_names: Set[str]) -> str:
    folded = {name.casefold(): name for name in available_names}
    alias = NAME_ALIASES.get(raw_name.casefold(), raw_name)
    return folded.get(alias.casefold(), "")


def extract_team_rosters(workbook, available_names: Set[str]) -> TeamSheetResult:
    """Read manager/team membership from the workbook's semantic fill colours."""
    if TEAM_SHEET_NAME not in workbook.sheetnames:
        raise ValueError(f"Required worksheet {TEAM_SHEET_NAME!r} not found")
    sheet = workbook[TEAM_SHEET_NAME]
    colour_to_manager: Dict[str, str] = {}
    ignored_colours: Set[str] = set()
    for row in range(1, min(sheet.max_row, 20) + 1):
        manager = str(sheet.cell(row, 10).value or "").strip()
        colour = _colour_hex(sheet.cell(row, 9).fill.fgColor, workbook)
        if manager in MANAGER_TO_TEAM and colour:
            colour_to_manager[colour] = manager
        elif manager.casefold() == "doppelt" and colour:
            ignored_colours.add(colour)
    if set(MANAGER_TO_TEAM.values()) - {MANAGER_TO_TEAM[m] for m in colour_to_manager.values()}:
        raise ValueError(f"Incomplete manager colour legend: {colour_to_manager}")

    teams: Dict[str, List[str]] = {team: [] for team in TEAM_CONFIG}
    unresolved: List[str] = []
    for row in range(4, min(sheet.max_row, 47) + 1):
        for column in range(2, min(sheet.max_column, 7) + 1):
            cell = sheet.cell(row, column)
            raw_name = _clean_sheet_name(cell.value)
            if not raw_name:
                continue
            colour = _colour_hex(cell.fill.fgColor, workbook)
            if not colour or colour in ignored_colours or colour not in colour_to_manager:
                continue
            resolved = _resolve_name(raw_name, available_names)
            if not resolved:
                unresolved.append(f"{cell.coordinate}: {raw_name}")
                continue
            team = MANAGER_TO_TEAM[colour_to_manager[colour]]
            if resolved not in teams[team]:
                teams[team].append(resolved)
    if unresolved:
        raise ValueError("Unresolved team wrestler names: " + ", ".join(unresolved))
    if any(not members for members in teams.values()):
        raise ValueError(f"A team resolved to an empty roster: {teams}")
    return TeamSheetResult(sheet.title, colour_to_manager, teams)


def fetch_team_workbook(sheet_id: str, credentials_path: Path):
    """Download the source Office workbook via Drive API, preserving cell fills."""
    from google.oauth2.service_account import Credentials
    from googleapiclient.discovery import build
    from googleapiclient.http import MediaIoBaseDownload

    credentials = Credentials.from_service_account_file(
        str(credentials_path), scopes=["https://www.googleapis.com/auth/drive.readonly"]
    )
    drive = build("drive", "v3", credentials=credentials, cache_discovery=False)
    metadata = drive.files().get(
        fileId=sheet_id, fields="id,name,mimeType,modifiedTime"
    ).execute()
    expected = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    if metadata.get("mimeType") != expected:
        raise ValueError(f"Expected XLSX source, got {metadata.get('mimeType')}")
    buffer = BytesIO()
    downloader = MediaIoBaseDownload(buffer, drive.files().get_media(fileId=sheet_id))
    done = False
    while not done:
        _, done = downloader.next_chunk()
    buffer.seek(0)
    return load_workbook(buffer, data_only=False), metadata


def _wrestler_catalog(repo: Path) -> Tuple[Set[str], Mapping[str, dict]]:
    names: Set[str] = set()
    metadata: Dict[str, dict] = {}
    for path in (repo / "src/site/notes/Wrestler").glob("*.md"):
        frontmatter, _ = parse_frontmatter(path.read_text(encoding="utf-8"))
        name = str(frontmatter.get("title") or path.stem)
        names.add(name)
        metadata[name] = frontmatter
    return names, metadata


def _write_team_cache(
    repo: Path,
    result: TeamSheetResult,
    source_metadata: Mapping[str, object],
    sheet_id: str,
) -> None:
    cache_path = repo / CACHE_RELATIVE_PATH
    cache_path.parent.mkdir(parents=True, exist_ok=True)
    payload = {
        "sheet_id": sheet_id,
        "sheet_gid": DEFAULT_SHEET_GID,
        "tab_name": result.tab_name,
        "source_name": source_metadata.get("name"),
        "source_modified": source_metadata.get("modifiedTime"),
        "colour_to_manager": dict(result.colour_to_manager),
        "teams": dict(result.teams),
    }
    cache_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def _load_team_cache(repo: Path) -> TeamSheetResult:
    cache_path = repo / CACHE_RELATIVE_PATH
    if not cache_path.exists():
        raise RuntimeError("Team sheet unavailable and no validated team cache exists")
    payload = json.loads(cache_path.read_text(encoding="utf-8"))
    return TeamSheetResult(
        str(payload["tab_name"]), payload["colour_to_manager"], payload["teams"]
    )


def _slug_from_frontmatter(frontmatter: Mapping[str, object], name: str) -> str:
    permalink = str(frontmatter.get("permalink") or "")
    match = re.search(r"/wrestler/([^/]+)/?", permalink)
    if not match:
        raise ValueError(f"Wrestler {name!r} has no canonical permalink")
    return match.group(1)


def update_team_pages(repo: Path, result: TeamSheetResult) -> int:
    names, wrestler_metadata = _wrestler_catalog(repo)
    resolved_total = 0
    for team, members in result.teams.items():
        config = TEAM_CONFIG[team]
        cards = []
        for name in members:
            if name not in names:
                raise ValueError(f"Cached team wrestler no longer exists: {name}")
            frontmatter = wrestler_metadata[name]
            slug = _slug_from_frontmatter(frontmatter, name)
            image = str(frontmatter.get("img") or (
                "https://github.com/CptSpaulding1980/choke-slam-wrestling/releases/download/images/"
                + name.replace(" ", "_") + ".png"
            ))
            cards.append(
                '  <a href="/choke-slam-wrestling/wrestler/{}/" class="team-member-card">\n'
                '    <img src="{}" alt="{}">\n'
                '    <span>{}</span>\n'
                "  </a>".format(slug, html.escape(image, quote=True), html.escape(name), html.escape(name))
            )
        resolved_total += len(cards)
        source = (
            "---\n"
            + yaml.safe_dump(
                {
                    "dg-publish": True,
                    "permalink": f"/teams/{config['slug']}/",
                    "title": team,
                },
                sort_keys=False,
                allow_unicode=True,
            )
            + "---\n\n"
            + f"# {team}\n\n"
            + '<div class="team-profile">\n'
            + '  <a class="team-manager" href="/choke-slam-wrestling/manager/{}/">\n'.format(config["manager_slug"])
            + '    <img src="https://github.com/CptSpaulding1980/choke-slam-wrestling/releases/download/images/{}" alt="{}">\n'.format(config["manager_image"], html.escape(config["manager"]))
            + "    <span><small>Manager</small>{}</span>\n  </a>\n".format(html.escape(config["manager"]))
            + '  <img class="team-logo" src="https://github.com/CptSpaulding1980/choke-slam-wrestling/releases/download/images/{}" alt="{} Logo">\n'.format(config["logo"], html.escape(team))
            + "</div>\n\n"
            + "## Entrance Theme\n"
            + '<audio controls preload="none"><source src="https://github.com/CptSpaulding1980/choke-slam-wrestling/releases/download/audio/{}" type="audio/mpeg"></audio>\n\n'.format(config["audio"])
            + "## Aktuelles Team\n"
            + TEAM_START + "\n"
            + '<div class="team-roster-grid">\n'
            + "\n".join(cards)
            + "\n</div>\n"
            + TEAM_END + "\n"
        )
        (repo / "src/site/notes/Teams" / f"{team}.md").write_text(source, encoding="utf-8")
    return resolved_total


def _championship_assignments(repo: Path) -> Mapping[str, List[Tuple[str, str]]]:
    assignments: Dict[str, List[Tuple[str, str]]] = {}
    available_names, _ = _wrestler_catalog(repo)
    pattern = re.compile(r"\[\[Wrestler/([^|\]\\]+)\\?\|([^\]]+)\]\]")
    for path in sorted((repo / "src/site/notes/Championships").glob("*.md")):
        text = path.read_text(encoding="utf-8")
        frontmatter, body = parse_frontmatter(text)
        title = str(frontmatter.get("title") or path.stem)
        permalink = str(frontmatter.get("permalink") or f"/championships/{path.stem}/")
        match = re.search(
            r"^##\s+👑\s+Aktuelle Champions[^\n]*\n(?P<line>[^\n]+)", body, flags=re.MULTILINE
        )
        if not match:
            continue
        for wrestler_target, _label in pattern.findall(match.group("line")):
            target = wrestler_target.strip()
            name = _resolve_name(target, available_names)
            if not name:
                raise ValueError(f"Current champion target does not resolve: {target}")
            assignments.setdefault(name, []).append((title, permalink))
    return assignments


def update_champion_panels(repo: Path) -> int:
    assignments = _championship_assignments(repo)
    relation_count = sum(len(titles) for titles in assignments.values())
    marker_pattern = re.compile(
        rf"\n?{re.escape(CHAMPION_START)}.*?{re.escape(CHAMPION_END)}\n?",
        re.DOTALL,
    )
    for path in sorted((repo / "src/site/notes/Wrestler").glob("*.md")):
        text = path.read_text(encoding="utf-8")
        frontmatter, _ = parse_frontmatter(text)
        name = str(frontmatter.get("title") or path.stem)
        text = marker_pattern.sub("\n", text)
        titles = assignments.get(name, [])
        if titles:
            links = "".join(
                '<a class="champion-title-link" href="/choke-slam-wrestling{}">{}</a>'.format(
                    permalink, html.escape(title)
                )
                for title, permalink in titles
            )
            panel = (
                f"\n{CHAMPION_START}\n"
                '<aside class="champion-status" aria-label="Aktueller Championstatus">\n'
                '  <span class="champion-status-label">Aktueller Champion</span>\n'
                f'  <div class="champion-status-titles">{links}</div>\n'
                "</aside>\n"
                f"{CHAMPION_END}\n"
            )
            profile_end = text.find("</table>")
            if profile_end == -1:
                raise ValueError(f"Wrestler profile table not found in {path}")
            profile_end += len("</table>")
            text = text[:profile_end] + panel + text[profile_end:]
        path.write_text(text, encoding="utf-8")
    return relation_count


def _season_eight_videos(repo: Path) -> List[dict]:
    videos: List[dict] = []
    for path in sorted((repo / "src/site/notes/Events").glob("*.md")):
        text = path.read_text(encoding="utf-8")
        frontmatter, body = parse_frontmatter(text)
        title = str(frontmatter.get("title") or path.stem)
        season_match = re.search(r"S08E\d+", title, flags=re.IGNORECASE)
        if not season_match:
            continue
        urls = re.findall(r'<source\s+[^>]*src="([^"]+\.mp4)"', body, flags=re.IGNORECASE)
        for index, url in enumerate(dict.fromkeys(urls), 1):
            videos.append(
                {
                    "title": title,
                    "date": str(frontmatter.get("date") or ""),
                    "permalink": str(frontmatter.get("permalink") or ""),
                    "poster": str(frontmatter.get("image") or ""),
                    "url": url,
                    "type": "full" if "full_event" in url.casefold() else "clip",
                    "part": index if len(urls) > 1 else None,
                }
            )
    return videos


def update_season_eight_highlights(repo: Path) -> int:
    videos = _season_eight_videos(repo)
    if not videos:
        raise ValueError("No Season 8 event videos found; refusing to erase the highlights section")
    cards = []
    for video in videos:
        title = video["title"].replace("_", " — ", 1)
        if video["part"]:
            title += f" · Teil {video['part']}"
        poster = video["poster"] or "https://github.com/CptSpaulding1980/choke-slam-wrestling/releases/download/images/ChokeSlam.png"
        cards.append(
            '    <article class="hl-card" data-type="{}" data-season="s08">\n'
            '      <video controls preload="metadata" poster="{}"><source src="{}" type="video/mp4"></video>\n'
            '      <div class="info"><div class="title">{}</div><div class="meta">{}</div>'
            '<a href="/choke-slam-wrestling{}">→ Event Page</a></div>\n'
            "    </article>".format(
                video["type"],
                html.escape(poster, quote=True),
                html.escape(video["url"], quote=True),
                html.escape(title),
                html.escape(video["date"]),
                video["permalink"],
            )
        )
    block = (
        HIGHLIGHTS_START
        + '\n<section class="hl-season" data-season="s08">\n'
        + "  <h2>Season 8 — Videos</h2>\n  <div class=\"hl-grid\">\n"
        + "\n".join(cards)
        + "\n  </div>\n</section>\n"
        + HIGHLIGHTS_END
    )
    path = repo / "src/site/notes/highlights.md"
    text = path.read_text(encoding="utf-8")
    marker_pattern = re.compile(
        rf"{re.escape(HIGHLIGHTS_START)}.*?{re.escape(HIGHLIGHTS_END)}", re.DOTALL
    )
    if marker_pattern.search(text):
        text = marker_pattern.sub(block, text)
    else:
        insert_at = text.find('<div class="hl-season"')
        if insert_at == -1:
            raise ValueError("Highlights page has no season insertion point")
        text = text[:insert_at] + block + "\n\n" + text[insert_at:]
    if "filter('s08')" not in text:
        all_button = re.search(r"(<button[^>]+filter\('all'\)[^>]*>.*?</button>)", text)
        if not all_button:
            raise ValueError("Highlights filter bar not found")
        text = text[: all_button.end()] + "\n  <button onclick=\"filter('s08')\">Season 8</button>" + text[all_button.end() :]
    path.write_text(text, encoding="utf-8")
    return len(videos)


def validate_source_event_links(repo: Path) -> ValidationResult:
    event_targets = {
        path.relative_to(repo / "src/site/notes").with_suffix("").as_posix()
        for path in (repo / "src/site/notes/Events").glob("*.md")
    }
    checked = 0
    missing: List[str] = []
    pattern = re.compile(r"\[\[(Events/[^|\]]+?)(?:\\?\|[^\]]+)?\]\]")
    for directory in ("Championships", "Wrestler", "Statistiken"):
        for path in (repo / "src/site/notes" / directory).glob("*.md"):
            for target in pattern.findall(path.read_text(encoding="utf-8")):
                checked += 1
                clean_target = target.rstrip("\\")
                if clean_target not in event_targets:
                    missing.append(f"{path.relative_to(repo)} -> {clean_target}")
    return ValidationResult(checked, missing)


def validate_built_event_links(repo: Path) -> ValidationResult:
    prefix = "/choke-slam-wrestling/events/"
    checked = 0
    missing: List[str] = []
    href_pattern = re.compile(r'href="(/choke-slam-wrestling/events/[^"#?]*/?)"')
    for directory in ("championships", "wrestler", "statistiken"):
        for path in (repo / "docs" / directory).glob("**/*.html"):
            text = path.read_text(encoding="utf-8")
            for href in href_pattern.findall(text):
                checked += 1
                relative = href[len(prefix) :].strip("/")
                target = repo / "docs/events" / relative / "index.html" if relative else repo / "docs/events/index.html"
                if not target.exists():
                    missing.append(f"{path.relative_to(repo)} -> {href}")
    return ValidationResult(checked, missing)


def validate_built_team_pages(repo: Path) -> ValidationResult:
    """Compare every cached roster member with the generated team HTML."""
    teams = _load_team_cache(repo)
    _, wrestler_metadata = _wrestler_catalog(repo)
    checked = 0
    missing: List[str] = []
    for team, members in teams.teams.items():
        config = TEAM_CONFIG.get(team)
        if not config:
            missing.append(f"Unknown cached team: {team}")
            continue
        page = repo / "docs/teams" / config["slug"] / "index.html"
        if not page.exists():
            missing.append(f"Missing built team page: {page.relative_to(repo)}")
            continue
        text = page.read_text(encoding="utf-8")
        rendered_cards = text.count('class="team-member-card"')
        if rendered_cards != len(members):
            missing.append(
                f"{page.relative_to(repo)} has {rendered_cards} cards; expected {len(members)}"
            )
        for name in members:
            checked += 1
            metadata = wrestler_metadata.get(name)
            if not metadata:
                missing.append(f"Cached wrestler missing from catalog: {name}")
                continue
            slug = _slug_from_frontmatter(metadata, name)
            expected = f'href="/choke-slam-wrestling/wrestler/{slug}/"'
            if expected not in text or f">{html.escape(name)}</span>" not in text:
                missing.append(f"{page.relative_to(repo)} missing roster member {name}")
    return ValidationResult(checked, missing)


def capture_event_media(repo: Path) -> Mapping[str, str]:
    snapshot: Dict[str, str] = {}
    for path in (repo / "src/site/notes/Events").glob("*.md"):
        text = path.read_text(encoding="utf-8")
        _, body = parse_frontmatter(text)
        matches_at = body.find("## Matches")
        if matches_at == -1:
            continue
        before_matches = body[:matches_at]
        first_media = min(
            (position for position in (before_matches.find("<video"), before_matches.find("<img")) if position >= 0),
            default=-1,
        )
        if first_media >= 0 and ("<video" in before_matches[first_media:] or "_poster" in before_matches[first_media:]):
            snapshot[path.name] = before_matches[first_media:].strip()
    return snapshot


def restore_event_media(repo: Path, snapshot: Mapping[str, str]) -> int:
    restored = 0
    for filename, media in snapshot.items():
        if Path(filename).name != filename or not filename.endswith(".md"):
            raise ValueError(f"Unsafe event snapshot filename: {filename!r}")
        path = repo / "src/site/notes/Events" / filename
        if not path.exists():
            continue
        text = path.read_text(encoding="utf-8")
        matches_at = text.find("## Matches")
        if matches_at == -1:
            continue
        prefix = text[:matches_at].rstrip()
        if "<video" in prefix or "_poster" in prefix:
            first_media = min(
                (position for position in (prefix.find("<video"), prefix.find("<img")) if position >= 0),
                default=-1,
            )
            if first_media >= 0:
                prefix = prefix[:first_media].rstrip()
        path.write_text(prefix + "\n\n" + media + "\n\n" + text[matches_at:], encoding="utf-8")
        restored += 1
    return restored


def sync_all(repo: Path, credentials: Path, sheet_id: str) -> None:
    names, _ = _wrestler_catalog(repo)
    try:
        workbook, metadata = fetch_team_workbook(sheet_id, credentials)
    except Exception as error:
        print(f"WARNING: Team sheet fetch failed ({error}); using validated cache", file=sys.stderr)
        teams = _load_team_cache(repo)
        source = "validated cache"
    else:
        teams = extract_team_rosters(workbook, names)
        _write_team_cache(repo, teams, metadata, sheet_id)
        source = f"live XLSX {metadata.get('name')} ({metadata.get('modifiedTime')})"
    team_count = update_team_pages(repo, teams)
    champion_count = update_champion_panels(repo)
    video_count = update_season_eight_highlights(repo)
    source_links = validate_source_event_links(repo)
    if source_links.missing:
        raise RuntimeError("Missing source event links:\n" + "\n".join(source_links.missing))
    print(f"Team source: {source}; tab={teams.tab_name!r}")
    print("Team colour mapping: " + ", ".join(f"{colour} -> {manager} -> {MANAGER_TO_TEAM[manager]}" for colour, manager in teams.colour_to_manager.items()))
    for team, members in teams.teams.items():
        print(f"Team roster: {team} ({len(members)}): {', '.join(members)}")
    print(f"Synced {team_count} team members, {champion_count} champion-title assignments, {video_count} S08 videos")
    print(f"Validated {source_links.checked} source event links")


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("command", choices=("sync", "validate-built", "capture-media", "restore-media"))
    parser.add_argument("--repo", type=Path, default=Path.cwd())
    parser.add_argument("--credentials", type=Path, default=DEFAULT_CREDENTIALS)
    parser.add_argument("--sheet-id", default=DEFAULT_SHEET_ID)
    parser.add_argument("--snapshot", type=Path)
    args = parser.parse_args(argv)
    repo = args.repo.resolve()
    if args.command == "sync":
        sync_all(repo, args.credentials, args.sheet_id)
    elif args.command == "validate-built":
        event_result = validate_built_event_links(repo)
        team_result = validate_built_team_pages(repo)
        print(f"Validated {event_result.checked} built event links")
        print(f"Validated {team_result.checked} built team memberships")
        missing = [*event_result.missing, *team_result.missing]
        if missing:
            print("\n".join(missing), file=sys.stderr)
            return 1
    elif args.command == "capture-media":
        if not args.snapshot:
            parser.error("capture-media requires --snapshot")
        args.snapshot.write_text(json.dumps(capture_event_media(repo), ensure_ascii=False), encoding="utf-8")
    elif args.command == "restore-media":
        if not args.snapshot or not args.snapshot.exists():
            parser.error("restore-media requires an existing --snapshot")
        restored = restore_event_media(repo, json.loads(args.snapshot.read_text(encoding="utf-8")))
        print(f"Restored media for {restored} events")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
