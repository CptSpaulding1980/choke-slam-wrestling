import importlib.util
import json
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook
from openpyxl.styles import PatternFill


MODULE_PATH = Path(__file__).parents[1] / "scripts" / "csw_site_sync.py"
SPEC = importlib.util.spec_from_file_location("csw_site_sync", MODULE_PATH)
site_sync = importlib.util.module_from_spec(SPEC)
sys.modules[SPEC.name] = site_sync
SPEC.loader.exec_module(site_sync)


class TeamSheetTests(unittest.TestCase):
    def test_extracts_only_manager_colours_and_resolves_names(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Choke Slam Wrestling S8"
        colours = {
            "Pascal": "FFFF9900",
            "Phillipp": "FFEAD1DC",
            "Hendrik": "FF61D836",
            "Doppelt": "FFCFE2F3",
        }
        for row, (manager, colour) in enumerate(colours.items(), 4):
            sheet.cell(row, 9).fill = PatternFill("solid", fgColor=colour)
            sheet.cell(row, 10).value = manager
        entries = [
            (4, 2, "Kazuchika Okada 1)", "Pascal"),
            (5, 2, "Mankind", "Phillipp"),
            (6, 2, "Steve Austin", "Hendrik"),
            (7, 2, "Ignored Duplicate", "Doppelt"),
        ]
        for row, column, name, manager in entries:
            sheet.cell(row, column).value = name
            sheet.cell(row, column).fill = PatternFill("solid", fgColor=colours[manager])

        result = site_sync.extract_team_rosters(
            workbook,
            available_names={"Kazuchika Okada", "Mick Foley", "Steve Austin"},
        )

        self.assertEqual(result.tab_name, "Choke Slam Wrestling S8")
        self.assertEqual(result.colour_to_manager["#FF9900"], "Pascal")
        self.assertEqual(result.colour_to_manager["#EAD1DC"], "Phillipp")
        self.assertEqual(result.colour_to_manager["#61D836"], "Hendrik")
        self.assertEqual(
            result.teams,
            {
                "Saint Rebel Radicalz": ["Kazuchika Okada"],
                "Militanter Mummenschanz": ["Mick Foley"],
                "Sweet 'n Sour Elite": ["Steve Austin"],
            },
        )

    def test_unresolved_wrestler_is_an_error(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Choke Slam Wrestling S8"
        for row, (manager, colour) in enumerate(
            (("Pascal", "FFFF9900"), ("Phillipp", "FFEAD1DC"), ("Hendrik", "FF61D836")),
            4,
        ):
            sheet.cell(row, 9).fill = PatternFill("solid", fgColor=colour)
            sheet.cell(row, 10).value = manager
        sheet["B4"] = "Unknown Wrestler"
        sheet["B4"].fill = PatternFill("solid", fgColor="FFFF9900")

        with self.assertRaisesRegex(ValueError, "Unknown Wrestler"):
            site_sync.extract_team_rosters(workbook, available_names={"Known Wrestler"})

    def test_invalid_live_workbook_does_not_fall_back_to_stale_cache(self):
        workbook = Workbook()
        sheet = workbook.worksheets[0]
        sheet.title = "Wrong sheet"
        with tempfile.TemporaryDirectory() as directory:
            repo = Path(directory)
            (repo / "src/site/notes/Wrestler").mkdir(parents=True)
            with patch.object(
                site_sync,
                "fetch_team_workbook",
                return_value=(workbook, {"name": "invalid.xlsx"}),
            ), patch.object(site_sync, "_load_team_cache") as cache_loader:
                with self.assertRaisesRegex(ValueError, "Required worksheet"):
                    site_sync.sync_all(repo, Path("credentials.json"), "custom-sheet-id")
            cache_loader.assert_not_called()

    def test_team_cache_records_requested_sheet_id(self):
        result = site_sync.TeamSheetResult(
            "Choke Slam Wrestling S8",
            {"#FF9900": "Pascal"},
            {"Saint Rebel Radicalz": ["Alice"]},
        )
        with tempfile.TemporaryDirectory() as directory:
            repo = Path(directory)
            site_sync._write_team_cache(repo, result, {}, "custom-sheet-id")
            payload = json.loads((repo / "data/team-rosters.json").read_text(encoding="utf-8"))
        self.assertEqual(payload["sheet_id"], "custom-sheet-id")


class GeneratedContentTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.repo = Path(self.temp_dir.name)
        (self.repo / "src/site/notes/Championships").mkdir(parents=True)
        (self.repo / "src/site/notes/Wrestler").mkdir(parents=True)
        (self.repo / "src/site/notes/Events").mkdir(parents=True)
        (self.repo / "src/site/notes/Teams").mkdir(parents=True)

    def tearDown(self):
        self.temp_dir.cleanup()

    def _write_wrestler(self, name):
        slug = name.lower().replace(" ", "-")
        (self.repo / f"src/site/notes/Wrestler/{name}.md").write_text(
            f"---\ntitle: {name}\npermalink: /wrestler/{slug}/\n---\n\n"
            f"# {name}\n\n<table><tr><td>Profile</td></tr></table>\n\n"
            "## Karriere-Statistiken\n",
            encoding="utf-8",
        )

    def test_champion_panels_support_multiple_titles_and_non_champion(self):
        for name in ("Alice", "Bob", "Carol"):
            self._write_wrestler(name)
        championships = self.repo / "src/site/notes/Championships"
        championships.joinpath("world.md").write_text(
            "---\ntitle: World Championship\npermalink: /championships/world/\n---\n\n"
            "## 👑 Aktuelle Champions (Singles)\n**[[Wrestler/Alice\\|Alice]]**\n",
            encoding="utf-8",
        )
        championships.joinpath("trios.md").write_text(
            "---\ntitle: Trios Championship\npermalink: /championships/trios/\n---\n\n"
            "## 👑 Aktuelle Champions (Trios)\n"
            "**[[Wrestler/Alice\\|Alice]] & [[Wrestler/Bob\\|Bob]]**\n",
            encoding="utf-8",
        )

        count = site_sync.update_champion_panels(self.repo)

        self.assertEqual(count, 3)
        alice = championships.parents[0].joinpath("Wrestler/Alice.md").read_text(encoding="utf-8")
        bob = championships.parents[0].joinpath("Wrestler/Bob.md").read_text(encoding="utf-8")
        carol = championships.parents[0].joinpath("Wrestler/Carol.md").read_text(encoding="utf-8")
        self.assertIn("champion-status", alice)
        self.assertIn("World Championship", alice)
        self.assertIn("Trios Championship", alice)
        self.assertIn("/championships/world/", alice)
        self.assertIn("champion-status", bob)
        self.assertNotIn("champion-status", carol)

    def test_champion_panels_use_canonical_target_not_display_alias(self):
        self._write_wrestler("Alice")
        self.repo.joinpath("src/site/notes/Championships/world.md").write_text(
            "---\ntitle: World Championship\npermalink: /championships/world/\n---\n\n"
            "## 👑 Aktuelle Champions (Singles)\n**[[Wrestler/Alice\\|The Ace]]**\n",
            encoding="utf-8",
        )

        site_sync.update_champion_panels(self.repo)

        alice = self.repo.joinpath("src/site/notes/Wrestler/Alice.md").read_text(encoding="utf-8")
        self.assertIn("World Championship", alice)

    def test_highlights_discovers_all_season_eight_event_videos(self):
        highlights = self.repo / "src/site/notes/highlights.md"
        highlights.write_text(
            "---\npermalink: /highlights/\n---\n\n"
            '<div class="hl-filter"><button onclick="filter(\'all\')">All</button></div>\n'
            '<div class="hl-season" data-season="s07"></div>\n',
            encoding="utf-8",
        )
        event = self.repo / "src/site/notes/Events/2026-09-14 - S08E02_All In.md"
        event.write_text(
            "---\ntitle: S08E02_All In\ndate: '2026-09-14'\n"
            "permalink: /events/2026-09-14-s08e02-all-in/\n---\n\n"
            '<source src="https://example.test/part-1_full_event.mp4" type="video/mp4">\n'
            '<source src="https://example.test/part-2_full_event.mp4" type="video/mp4">\n',
            encoding="utf-8",
        )

        count = site_sync.update_season_eight_highlights(self.repo)
        rendered = highlights.read_text(encoding="utf-8")

        self.assertEqual(count, 2)
        self.assertIn("Season 8", rendered)
        self.assertIn("part-1_full_event.mp4", rendered)
        self.assertIn("part-2_full_event.mp4", rendered)
        self.assertIn("/events/2026-09-14-s08e02-all-in/", rendered)

    def test_event_link_validator_handles_escaped_pipes_and_special_names(self):
        event = self.repo / "src/site/notes/Events/2024-01-01 - Show: Don't_Stop!.md"
        event.write_text(
            "---\npermalink: /events/2024-01-01-show-dont-stop/\n---\n",
            encoding="utf-8",
        )
        championship = self.repo / "src/site/notes/Championships/world.md"
        championship.write_text(
            "<table><tr><td>[[Events/2024-01-01 - Show: Don't_Stop!\\|Event]]</td></tr></table>",
            encoding="utf-8",
        )

        result = site_sync.validate_source_event_links(self.repo)

        self.assertEqual(result.checked, 1)
        self.assertEqual(result.missing, [])

    def test_event_media_survives_legacy_regeneration(self):
        event = self.repo / "src/site/notes/Events/2026-09-14 - S08E02_All In.md"
        event.write_text(
            "---\ntitle: S08E02_All In\n---\n\nHeader\n\n"
            '<img src="event_poster.png">\n<video><source src="part_full_event.mp4"></video>\n\n'
            "## Matches\nOld matches\n",
            encoding="utf-8",
        )
        snapshot = site_sync.capture_event_media(self.repo)
        event.write_text(
            "---\ntitle: S08E02_All In\n---\n\nRegenerated header\n\n"
            "## Matches\nNew matches\n",
            encoding="utf-8",
        )

        restored = site_sync.restore_event_media(self.repo, snapshot)
        text = event.read_text(encoding="utf-8")

        self.assertEqual(restored, 1)
        self.assertIn("event_poster.png", text)
        self.assertIn("part_full_event.mp4", text)
        self.assertIn("New matches", text)

    def test_event_media_rejects_snapshot_path_traversal(self):
        with self.assertRaisesRegex(ValueError, "Unsafe event snapshot filename"):
            site_sync.restore_event_media(self.repo, {"../../outside.md": "<video></video>"})

    def test_built_team_validator_compares_cache_to_rendered_cards(self):
        members = {
            "Militanter Mummenschanz": ["Alice"],
            "Saint Rebel Radicalz": ["Bob"],
            "Sweet 'n Sour Elite": ["Carol"],
        }
        for name in ("Alice", "Bob", "Carol"):
            self._write_wrestler(name)
        cache = self.repo / "data/team-rosters.json"
        cache.parent.mkdir(parents=True)
        cache.write_text(
            json.dumps(
                {
                    "tab_name": "Choke Slam Wrestling S8",
                    "colour_to_manager": {},
                    "teams": members,
                }
            ),
            encoding="utf-8",
        )
        for team, names in members.items():
            config = site_sync.TEAM_CONFIG[team]
            page = self.repo / "docs/teams" / config["slug"] / "index.html"
            page.parent.mkdir(parents=True)
            name = names[0]
            slug = name.lower()
            page.write_text(
                f'<a class="team-member-card" href="/choke-slam-wrestling/wrestler/{slug}/">'
                f"<span>{name}</span></a>",
                encoding="utf-8",
            )

        result = site_sync.validate_built_team_pages(self.repo)

        self.assertEqual(result.checked, 3)
        self.assertEqual(result.missing, [])


if __name__ == "__main__":
    unittest.main()
